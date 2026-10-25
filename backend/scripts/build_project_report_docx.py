"""Generate docs/PROJECT_REPORT.docx for the Insurance Claims Intelligence project."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "PROJECT_REPORT.docx"


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def main() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading("Insurance Claims Intelligence — Project Report", 0)
    p = doc.add_paragraph()
    p.add_run("Week 2 (Agentic Hybrid RAG) + Week 3 (Single-Agent Claim Process)").italic = True
    doc.add_paragraph(
        "This report summarizes what was built, which datasets power the demo, how prompts "
        "evolved during iterative (“vibe”) coding, the main iterations tried, and learnings "
        "from the workflow."
    )

    # ----- Overview -----
    doc.add_heading("1. Project overview", 1)
    doc.add_paragraph(
        "Insurance Claims Intelligence is a local demo application that helps a human adjuster "
        "investigate synthetic auto claims. An AI agent can look up claim and policy facts, "
        "retrieve policy language, analyze repair estimates, optionally search the public web "
        "(weather / market context), and draft an investigation report. The system never "
        "approves, denies, settles, or pays a claim."
    )

    doc.add_heading("1.1 Week 2 — Agentic Hybrid RAG", 2)
    add_bullets(
        doc,
        [
            "FastAPI backend + React / React Flow frontend",
            "LangGraph tool-calling agent on Nebius Token Factory chat models",
            "Hybrid policy retrieval: local BM25 + Nebius embeddings + Pinecone + Reciprocal Rank Fusion (RRF)",
            "Deterministic tools for claim intake, policy dates, coverage flags, estimate math, missing docs, similar claims, portfolio stats",
            "Live SSE event stream driving a workflow canvas and investigation report UI",
            "Safety: human_review_required always; illustrative financials only",
        ],
    )

    doc.add_heading("1.2 Week 3 — Single-agent claim process", 2)
    add_bullets(
        doc,
        [
            "Stages: Filing → Classification → Investigation → Report → Human Review",
            "New tools: classify_claim, list_claim_documents, parse/build/compare estimates, Keenable search_web, market cost research, request_human_handoff",
            "Preliminary (desk) vs facility (shop) estimate Markdown per claim",
            "Keenable web search for weather and illustrative parts/labor market context",
            "Adjuster Desk UI (ask + run + review) and Live Workflow debug canvas (shared run state; click node → input / output / logs)",
            "Still single-agent (not multi-agent); tool retry + human handoff on repeated failures",
            "Report money figures verified in code (Decimal) so the model cannot invent estimate deltas",
        ],
    )

    doc.add_heading("1.3 Stack", 2)
    add_bullets(
        doc,
        [
            "Backend: Python, FastAPI, LangGraph, uv",
            "Frontend: React, TypeScript, Vite, React Flow",
            "Cloud: Nebius (chat + embeddings), Pinecone (dense vectors), Keenable (web search)",
            "Local: synthetic JSON/Markdown datasets, BM25 index",
        ],
    )

    # ----- Datasets -----
    doc.add_heading("2. Datasets used", 1)
    doc.add_paragraph(
        "All business data is synthetic and generated under backend/data/ (scripts/generate_data.py). "
        "Approximate current sizes:"
    )
    add_bullets(
        doc,
        [
            "5 policies (JSON + Markdown under data/policies/)",
            "15 claims (claims.json) with scenario tags (collision, comprehensive/vandalism, hail, expired policy, duplicate lines, etc.)",
            "~70 claim document metadata rows (claim_documents.json)",
            "~58 repair line items (repair_line_items.json / .csv)",
            "15 claim folders with FNOL.md, Preliminary_Estimate.md, Repair_Facility_Estimate.md",
            "Required-document checklists (required_documents.json)",
            "~20 evaluation questions (evaluation_questions.json)",
            "Policy clause library snippets used in generation / RAG corpus",
            "BM25 index under data/indexes/ after ingest; Pinecone namespace for dense chunks",
        ],
    )
    doc.add_paragraph(
        "Claim.estimate_total is aligned with the facility line-item sum when data is regenerated, "
        "so desk/report totals stay consistent with Repair_Facility_Estimate.md."
    )

    # ----- Prompts -----
    doc.add_heading("3. Prompts used during vibe coding", 1)
    doc.add_paragraph(
        "Two layers of prompting shaped the product: (A) conversational prompts used while "
        "building with the coding assistant, and (B) the runtime agent system / user prompts "
        "shipped in the app."
    )

    doc.add_heading("3.1 Build-time prompts (examples of intent)", 2)
    doc.add_paragraph(
        "These are representative of the iterative asks used while developing (paraphrased):"
    )
    add_bullets(
        doc,
        [
            "Update README / Windows setup; produce a Word project report (not Markdown-only).",
            "Extend Week 2 into Week 3: full claim process Filing→Classify→Investigate→Report→Human Review; use web search (Tavily then Keenable) for weather and market parts/labor.",
            "Prefer a single agent, not multi-agent.",
            "Make estimates more realistic: preliminary vs facility; show change after vehicle at shop.",
            "Show the process in the UI (pipeline + canvas nodes/edges); add an Adjuster Desk tab.",
            "Live Workflow: only show nodes; on click show input, output, and logs; remove other chrome.",
            "Runs started on Adjuster Desk must continue on Live Workflow (shared state).",
            "No Run on Live Workflow — debug mirror only; Adjuster Desk needs a free-form question box.",
            "Every comprehensive run must search weather / required web info.",
            "Fix recursion limit / wrong estimate math in summaries; clarify illustrative financials vs market research.",
            "Create architecture doc and short demo recording scripts (Week 3 clips only).",
        ],
    )

    doc.add_heading("3.2 Runtime agent system prompt (shipped)", 2)
    doc.add_paragraph(
        "Source: backend/app/agent/prompts.py — SYSTEM_PROMPT. Core instructions:"
    )
    add_bullets(
        doc,
        [
            "Role: Insurance Claims Process Agent (single-agent, multi-stage).",
            "Never invent policy text, dates, deductibles, counts, URLs, or dollar amounts.",
            "Stages in order for claim-scoped runs: Filing → Classification → Investigation → Report → Human review.",
            "Investigation: map question types to tools; on comprehensive runs require search_web (weather + public context) and market cost research.",
            "Estimate rules: prelim may change at facility; copy tool deltas verbatim; Keenable $ are illustrative; hand off if tools fail.",
            "Portfolio mode: query_portfolio_data then report with claim_id=portfolio.",
            "Answer style: claim_summary answers the question in 1–3 sentences; structured sections for facts/evidence; human_review_required; never approve/deny/settle/pay.",
            "Prefer batching independent tool calls on comprehensive runs to stay within recursion limits.",
        ],
    )

    doc.add_heading("3.3 Runtime user-message templates", 2)
    doc.add_paragraph(
        "Source: backend/app/agent/graph.py — built per run from claim_id + question:"
    )
    add_bullets(
        doc,
        [
            "Process / comprehensive mode: ordered checklist (filing tools → classify → estimates → REQUIRED search_web + research_market_repair_costs → generate_investigation_report).",
            "Narrow mode: answer only the question; call needed tools; finish with report.",
            "Portfolio mode: no claim; use query_portfolio_data; report claim_id=portfolio.",
            "UI default Adjuster Desk question (paraphrase): run a comprehensive full claim process including weather and required web research.",
        ],
    )

    # ----- Iterations -----
    doc.add_heading("4. Iterations tried", 1)

    doc.add_heading("4.1 Week 2 iterations", 2)
    add_bullets(
        doc,
        [
            "Hybrid RAG wiring (BM25 + dense + RRF) with live retrieval nodes on the canvas.",
            "Tool-first agent with force_report when the model forgot generate_investigation_report.",
            "Citation backfill from tool messages when the model omitted retrieved_evidence.",
            "Portfolio vs claim-scoped modes and clearer claim_summary instructions.",
            "Windows setup hardening (uv, PowerShell scripts, env examples, truststore / IDE notes).",
        ],
    )

    doc.add_heading("4.2 Week 3 product iterations", 2)
    add_bullets(
        doc,
        [
            "Started from assignment goals (agentic process + web search + estimates); briefly considered multi-agent, then locked to single-agent with stages.",
            "Web search provider: Tavily planned → switched to Keenable (API key + client + health light).",
            "Estimate docs: generate preliminary + facility Markdown; compare tool; sync claim.estimate_total to line sums.",
            "UI: ProcessPipeline stages; then Adjuster Desk as default human workspace; Live Workflow simplified to nodes + I/O/logs.",
            "Shared InvestigationProvider so desk and workflow share one SSE run.",
            "Comprehensive runs: prompts require weather + market research; raised MAX_TOOL_CALLS / recursion_limit / timeout after GRAPH_RECURSION_LIMIT failures.",
            "Wrong LLM math on estimate deltas → overwrite illustrative financials / estimate_analysis in generate_investigation_report and scrub invented $ sentences from claim_summary.",
        ],
    )

    doc.add_heading("4.3 Iteration that did not stick", 2)
    add_bullets(
        doc,
        [
            "Multi-agent design — rejected for Week 3 scope; one agent with more tools was clearer to demo and operate.",
            "Tavily integration — replaced by Keenable before shipping.",
            "Heavy Live Workflow chrome (sidebar, bottom logs, minimap) — removed so the canvas stays a debug view.",
            "Trusting LLM prose for dollar comparisons — replaced with code verification.",
        ],
    )

    # ----- Learnings -----
    doc.add_heading("5. Learnings and observations", 1)
    add_bullets(
        doc,
        [
            "Stage prompts + tool lists work better than hoping the model invents a process; comprehensive mode must be explicit about required tools (weather/market).",
            "LangGraph recursion limits bite when the model calls tools one-at-a-time; batching guidance and higher limits are both needed for full process runs.",
            "LLMs invent arithmetic even when tools return correct totals — money belonging in tools/code, not free text.",
            "Illustrative financials (facility − deductible) must stay separate from market research (web context); operators confuse them unless the UI and report label clearly.",
            "Adjuster-facing UI and engineer debug UI should share state but not the same layout; splitting Desk vs Workflow improved demos.",
            "Synthetic data drifts (estimate_total vs line sums) silently poison demos — generate_data should keep them aligned.",
            "External search (Keenable) needs health checks, retries, and a no-invention rule when search fails.",
            "Human boundary language must be repeated in prompts, report disclaimer, and UI — otherwise audiences assume the agent “pays” the claim.",
            "Vibe coding works well for vertical slices (tool → API → canvas node → desk panel) with frequent demo runs; documentation (README, architecture, demo scripts) needs to be rewritten when files get wiped or scope shifts.",
        ],
    )

    doc.add_heading("6. Safety summary", 1)
    add_bullets(
        doc,
        [
            "Policy-date membership ≠ legal coverage determination.",
            "Coverage listed ≠ coverage decision.",
            "Estimate nets and market prices are illustrative only.",
            "Duplicates flagged, never auto-removed.",
            "No approve / deny / settle / pay — licensed human adjuster decides.",
        ],
    )

    doc.add_heading("7. Related documents", 1)
    add_bullets(
        doc,
        [
            "README.md — setup and Week 3 product overview",
            "docs/ARCHITECTURE.md — system diagrams and component map",
            "docs/WEEK3_AGENTIC.md — Week 3 stages, tools, demo claims",
            "docs/WEEK3_DEMO_SCRIPTS.docx — short recording clips for Week 3 only",
            "SETUP_WINDOWS.md — Windows troubleshooting",
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
